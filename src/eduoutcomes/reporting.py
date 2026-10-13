"""Registro y exportación de tablas y figuras (DOCX de tres líneas, LaTeX ``sn-jnl``, Excel, PDF/EPS/TIFF/PNG)."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image

from .config import Config
from .paths import Paths
from .style import COL_MM, MAX_H_MM, MM, align_panel_labels, min_font

log = logging.getLogger(__name__)


@dataclass
class Reporter:
    cfg: Config
    paths: Paths
    tables: list[dict] = field(default_factory=list)
    captions: dict[str, str] = field(default_factory=dict)
    figure_qa: list[dict] = field(default_factory=list)

    # ------------------------------------------------------------------------------- tablas
    def register_table(self, df: pd.DataFrame, key: str, number: str, title: str, note: str = "",
                       supplementary: bool = False, landscape: bool = False, group_col: str | None = None,
                       bold=None) -> None:
        """Registra una tabla para DOCX/LaTeX/XLSX y la guarda en CSV.

        group_col: columna cuyos valores se convierten en filas de grupo que abarcan toda la tabla.
        bold: conjunto de (fila, columna) que se resaltan en negrita.
        """
        folder = self.paths.supplementary if supplementary else self.paths.tables
        df = df.reset_index(drop=True)
        df.to_csv(folder / f"{key}.csv", index=False, encoding="utf-8-sig")
        self.tables = [t for t in self.tables if t["key"] != key]   # re-ejecución idempotente
        self.tables.append(dict(df=df.copy(), key=key, number=number, title=title, note=note,
                                supplementary=supplementary, landscape=landscape, group_col=group_col,
                                bold=set(bold or ())))
        log.info("Table %s registered — %s", number, title)
        with pd.option_context("display.max_columns", None, "display.width", 220, "display.max_colwidth", 60):
            log.debug("\n%s", df.to_string(index=False))

    # ------------------------------------------------------------------------------ figuras
    def save_figure(self, fig, stem: str, caption: str, supplementary: bool = False) -> None:
        """Guarda PDF/EPS vectoriales (fuentes incrustadas), TIFF RGB LZW y PNG; registra el control de calidad."""
        align_panel_labels(fig)
        folder = self.paths.supplementary if supplementary else self.paths.figures
        w_mm, h_mm = (fig.get_size_inches() / MM).round(1)
        for ext in self.cfg.fig_formats:
            path = folder / f"{stem}.{ext}"
            if ext == "tiff":
                buf = io.BytesIO()
                fig.savefig(buf, format="png", dpi=self.cfg.tiff_dpi, facecolor="white")
                buf.seek(0)
                Image.open(buf).convert("RGB").save(path, compression="tiff_lzw",
                                                    dpi=(self.cfg.tiff_dpi, self.cfg.tiff_dpi))
            elif ext == "png":
                fig.savefig(path, dpi=self.cfg.png_dpi, facecolor="white")
            else:
                fig.savefig(path, format=ext, facecolor="white")
        mf = min_font(fig)
        self.figure_qa = [q for q in self.figure_qa if q["figure"] != stem]
        self.figure_qa.append({
            "figure": stem, "width_mm": float(w_mm), "height_mm": float(h_mm), "min_font_pt": mf,
            "width_ok": float(w_mm) in [float(v) for v in COL_MM.values()], "height_ok": h_mm <= MAX_H_MM,
            "font_ok": mf >= self.cfg.min_font_pt - 1e-9, "formats": ",".join(self.cfg.fig_formats)})
        self.captions[stem] = caption
        log.info("Figure saved — %s (%.0f × %.0f mm)", stem, w_mm, h_mm)
        if self.cfg.show_figures:
            plt.show()
        plt.close(fig)

    # --------------------------------------------------------------------------- exportación
    def export_all(self, extra_sheets: dict[str, pd.DataFrame] | None = None) -> None:
        main_t = sorted([t for t in self.tables if not t["supplementary"]], key=lambda t: int(t["number"]))
        supp_t = sorted([t for t in self.tables if t["supplementary"] and t["key"] != "TableS4_hyperparameters"],
                        key=lambda t: int(t["number"][1:]))
        write_docx(main_t, self.paths.tables / "Tables_Springer.docx", "Tables")
        write_docx(supp_t, self.paths.supplementary / "Supplementary_Tables.docx", "Supplementary tables")
        write_latex(main_t, self.paths.tables / "tables_sn-jnl.tex")
        write_latex(supp_t, self.paths.supplementary / "supplementary_tables_sn-jnl.tex")
        write_excel(self.tables, self.paths.root / "results_all.xlsx", extra_sheets or {})
        (self.paths.root / "figure_captions.txt").write_text("\n\n".join(self.captions.values()), encoding="utf-8")


# =============================================================================================
# Utilidades comunes a DOCX y LaTeX
# =============================================================================================
_NUMERIC = re.compile(r"^[\s<>≤≥]*[\u2212\-+]?\d")


def _is_numeric_col(s: pd.Series) -> bool:
    v = [x for x in s.astype(str) if x not in ("", "NE", "—")]
    return bool(v) and np.mean([bool(_NUMERIC.match(x)) for x in v]) >= 0.8


def _layout(t: dict):
    """Devuelve (df a mostrar, filas de grupo {posición: rótulo}, columnas numéricas, celdas en negrita)."""
    df_ = t["df"].astype(str).replace({"nan": "", "None": "", "<NA>": ""})
    gcol, groups = t.get("group_col"), {}
    if gcol and gcol in df_.columns:
        prev = None
        for i, g in enumerate(df_[gcol]):
            if g != prev:
                groups[i] = g
                prev = g
        df_ = df_.drop(columns=gcol)
    bold = {(i, c) for i, c in t.get("bold", set()) if c in df_.columns}
    numeric = {c for j, c in enumerate(df_.columns) if j > 0 and _is_numeric_col(df_[c])}
    return df_, groups, numeric, bold


def _col_widths(df_: pd.DataFrame, usable: float, mm_per_char: float = 1.45, pad: float = 2.6) -> np.ndarray:
    """Anchos (mm) que evitan partir números (cada columna recibe al menos su palabra más larga)."""
    cols = list(df_.columns)
    full = np.array([max([len(v) for v in df_[c]] + [len(str(c)) * 0.6, 1]) for c in cols], float)
    word = np.array([max(len(w) for v in list(df_[c]) + [str(c)] for w in (str(v).split() or [""])) for c in cols],
                    float)
    need = full * mm_per_char + pad
    minw = np.minimum(word * mm_per_char * 1.1 + pad, need.max())
    need = np.maximum(need, minw)
    if need.sum() <= usable:
        return need * usable / need.sum()
    flex, deficit = need - minw, need.sum() - usable
    if flex.sum() >= deficit:
        return need - flex * deficit / flex.sum()
    return minw * usable / minw.sum()


# =============================================================================================
# DOCX (tablas de tres líneas, estilo Springer)
# =============================================================================================
def write_docx(tables: list[dict], path: Path, heading: str) -> None:
    try:
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt
    except ImportError:
        log.warning("python-docx not installed: DOCX export skipped")
        return

    def set_border(cell, **edges):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        for edge, val in edges.items():
            el = OxmlElement(f"w:{edge}")
            for k, v in val.items():
                el.set(qn(f"w:{k}"), str(v))
            borders.append(el)

    def tight(par, before=1, after=1):
        pf = par.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0

    def row_flag(row, tag):
        tr_pr = row._tr.get_or_add_trPr()
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "true")
        tr_pr.append(el)

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Mm(25))
    d.add_paragraph().add_run(heading).bold = True
    line = {"val": "single", "sz": 6, "color": "000000"}   # 0.75 pt
    thin = {"val": "single", "sz": 4, "color": "000000"}   # 0.5 pt
    none = {"val": "nil"}
    first = True
    for t in tables:
        land = t["landscape"]
        if not first or land:
            new = d.add_section()
            new.orientation = WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
            new.page_width, new.page_height = (Mm(297), Mm(210)) if land else (Mm(210), Mm(297))
            for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
                setattr(new, m, Mm(20 if land else 25))
        first = False
        df_, groups, numeric, bold = _layout(t)
        p = d.add_paragraph()
        tight(p, 0, 4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"Table {t['number']} ")
        r.bold, r.font.size = True, Pt(9)
        p.add_run(t["title"]).font.size = Pt(9)
        tb = d.add_table(rows=1, cols=df_.shape[1])
        tb.autofit = False
        row_flag(tb.rows[0], "tblHeader")
        for j, c in enumerate(df_.columns):
            cell = tb.rows[0].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(c))
            run.bold, run.font.size = True, Pt(8)
            tight(cell.paragraphs[0])
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c in numeric else WD_ALIGN_PARAGRAPH.LEFT
            set_border(cell, top=line, bottom=thin, left=none, right=none)
        n_rows = len(df_)
        for i, row in enumerate(df_.itertuples(index=False)):
            if i in groups and groups[i]:
                gr = tb.add_row()
                row_flag(gr, "cantSplit")
                mc = gr.cells[0].merge(gr.cells[-1])
                mc.text = ""
                rr = mc.paragraphs[0].add_run(groups[i])
                rr.italic, rr.font.size = True, Pt(8)
                tight(mc.paragraphs[0], 3, 1)
                set_border(mc, top=none, bottom=none, left=none, right=none)
            cells = tb.add_row().cells
            row_flag(tb.rows[-1], "cantSplit")
            last = i == n_rows - 1
            for j, v in enumerate(row):
                cells[j].text = ""
                run = cells[j].paragraphs[0].add_run(v)
                run.font.size = Pt(8)
                run.bold = (i, df_.columns[j]) in bold
                tight(cells[j].paragraphs[0])
                cells[j].paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.CENTER if df_.columns[j] in numeric
                                                    else WD_ALIGN_PARAGRAPH.LEFT)
                set_border(cells[j], bottom=line if last else none, left=none, right=none, top=none)
        wmm = _col_widths(df_, (297 - 40) if land else (210 - 50))
        for j, col in enumerate(tb.columns):
            col.width = Mm(wmm[j])
        for row in tb.rows:
            for j, cell in enumerate(row.cells):
                if j < len(wmm):
                    cell.width = Mm(wmm[j])
        if t["note"]:
            np_ = d.add_paragraph()
            tight(np_, 3, 0)
            np_.add_run(t["note"] if t["note"].endswith(".") else t["note"] + ".").font.size = Pt(8)
    d.save(path)
    log.info("Saved %s", path)


# =============================================================================================
# LaTeX (plantilla sn-jnl de Springer Nature)
# =============================================================================================
_TEX = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "_": r"\_", "#": r"\#", "$": r"\$", "{": r"\{",
        "}": r"\}", "^": r"\^{}", "~": r"\textasciitilde{}", "±": r"$\pm$", "≥": r"$\geq$", "≤": r"$\leq$",
        "χ": r"$\chi$", "²": r"$^{2}$", "−": r"$-$", "κ": r"$\kappa$", "Δ": r"$\Delta$", "×": r"$\times$",
        "α": r"$\alpha$", "τ": r"$\tau$", "ρ": r"$\rho$", "ε": r"$\varepsilon$", "Σ": r"$\Sigma$",
        "†": r"$^{\dagger}$", "‡": r"$^{\ddagger}$", "·": r"$\cdot$", "≈": r"$\approx$", "→": r"$\rightarrow$",
        "<": r"$<$", ">": r"$>$", "–": "--", "—": "---", "|": r"$|$", "’": "'", "'": "'"}


def latex_escape(s) -> str:
    return "".join(_TEX.get(ch, ch) for ch in str(s))


def write_latex(tables: list[dict], path: Path) -> None:
    out = ["% Tables for the Springer Nature LaTeX template (sn-jnl.cls). Requires booktabs; landscape tables use\n"
           "% sidewaystable (rotating package, loaded by sn-jnl). Bold marks the best value per column.\n"
           "\\providecommand{\\botrule}{\\bottomrule}\n\n"]
    for t in tables:
        df_, groups, numeric, bold = _layout(t)
        ncol = df_.shape[1]
        env = "sidewaystable" if t["landscape"] else "table"
        spec = "@{}" + "".join("c" if c in numeric else "l" for c in df_.columns) + "@{}"
        out.append(f"\\begin{{{env}}}[!htbp]\n\\caption{{{latex_escape(t['title'])}}}\\label{{tab:{t['key']}}}\n"
                   + ("\\footnotesize\n" if t["landscape"] or ncol > 6 else "\\small\n")
                   + f"\\begin{{tabular*}}{{\\textwidth}}{{@{{\\extracolsep\\fill}}{spec}}}\n\\toprule\n")
        out.append(" & ".join(latex_escape(c) for c in df_.columns) + r" \\" + "\n\\midrule\n")
        for i, row in enumerate(df_.itertuples(index=False)):
            if i in groups and groups[i]:
                if i > 0:
                    out.append("\\addlinespace[2pt]\n")
                out.append(f"\\multicolumn{{{ncol}}}{{@{{}}l}}{{\\textit{{{latex_escape(groups[i])}}}}} \\\\\n")
            cells = [(r"\textbf{" + latex_escape(v) + "}") if (i, df_.columns[j]) in bold else latex_escape(v)
                     for j, v in enumerate(row)]
            out.append(" & ".join(cells) + r" \\" + "\n")
        out.append("\\botrule\n\\end{tabular*}\n")
        if t["note"]:
            out.append(f"\\footnotetext{{{latex_escape(t['note'])}.}}\n")
        out.append(f"\\end{{{env}}}\n\n")
    Path(path).write_text("".join(out), encoding="utf-8")
    log.info("Saved %s", path)


# =============================================================================================
# Excel con formato
# =============================================================================================
def _xl_sheet(xw, df_: pd.DataFrame, name: str, title: str | None = None, note: str | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    start = 2 if title else 0
    df_.to_excel(xw, sheet_name=name, index=False, startrow=start)
    ws = xw.sheets[name]
    if title:
        ws.cell(1, 1, title).font = Font(bold=True)
    thin = Side(style="thin", color="000000")
    for j, c in enumerate(df_.columns, 1):
        h = ws.cell(start + 1, j)
        h.font = Font(bold=True)
        h.alignment = Alignment(wrap_text=True, vertical="center")
        h.border = Border(top=thin, bottom=thin)
        w = max([len(str(c))] + [len(str(v)) for v in df_[c].head(500)]) if len(df_) else len(str(c))
        ws.column_dimensions[get_column_letter(j)].width = min(max(8, w * 1.05 + 1), 60)
    ws.freeze_panes = ws.cell(start + 2, 1)
    if note:
        ws.cell(start + len(df_) + 3, 1, "Note. " + note).alignment = Alignment(wrap_text=False)


def write_excel(tables: list[dict], path: Path, extra: dict[str, pd.DataFrame]) -> None:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        log.warning("openpyxl not installed: Excel export skipped")
        return
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for t in tables:
            _xl_sheet(xw, t["df"], f"T{t['number']}"[:31], f"Table {t['number']}. {t['title']}", t["note"])
        for nm, d_ in extra.items():
            if d_ is not None:
                _xl_sheet(xw, d_, nm[:31])
    log.info("Saved %s", path)
